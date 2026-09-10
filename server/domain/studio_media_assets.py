"""创作域视频、音频与通用文件资产。"""

from __future__ import annotations

import hashlib
import mimetypes
import re
import uuid
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from domain.imports import decode_text
from domain.models import StudioMediaAsset
from domain.storage import get_storage

MEDIA_KINDS = frozenset({"video", "audio", "file"})
VIDEO_EXTENSIONS = frozenset({".mp4", ".webm", ".mov", ".m4v", ".avi", ".mkv"})
AUDIO_EXTENSIONS = frozenset({".mp3", ".wav", ".m4a", ".aac", ".ogg", ".flac"})


class StudioMediaAssetError(ValueError):
    pass


# ---------------------------------------------------------------- 附件正文抽取

# 单个附件最多喂多少字进提示词。给得太多的后果不是报错而是**规划变差**：
# 一份 80 页的规范塞满上下文，模型会盯着最后几段写提示词，把用户那句需求挤没。
ATTACHMENT_TEXT_CHARS = 4000

# 认得出正文的扩展名。二进制格式（docx/xlsx/pptx）不在里面——它们是 zip 包，
# 直接解码只会得到一堆乱码，那比"读不出"更糟：模型会把乱码当内容去用。
_TEXT_EXTS = frozenset(
    {
        ".txt", ".md", ".markdown", ".rst", ".log",
        ".json", ".yaml", ".yml", ".toml", ".ini", ".csv", ".tsv",
        ".html", ".htm", ".xml", ".srt", ".vtt",
        ".py", ".js", ".ts", ".tsx", ".jsx", ".css", ".scss", ".sql", ".sh",
    }
)


def _pdf_text(data: bytes, limit: int) -> str:
    from io import BytesIO

    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    out: list[str] = []
    total = 0
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if not text:
            continue
        out.append(text)
        total += len(text)
        if total >= limit:
            break
    return "\n\n".join(out)


def document_text(name: str, data: bytes, limit: int = ATTACHMENT_TEXT_CHARS) -> str:
    """尽力抽出附件正文，抽不出就返回空串。

    **抽不出不是错误**。用户带一个 sketch 文件进来是完全正当的，规划时只是
    没有正文可读而已；抛异常的话整次规划会因为一个附件失败，这个代价不成比例。
    调用方按空串处理：文件名照样告诉模型"用户带了这个"。
    """
    ext = ("." + name.rsplit(".", 1)[-1].lower()) if "." in name else ""
    try:
        if ext == ".pdf":
            text = _pdf_text(data, limit)
        elif ext in _TEXT_EXTS:
            text = decode_text(data)
        else:
            return ""
    except Exception:  # noqa: BLE001 —— 见 docstring：抽不出不该拖垮整次规划
        return ""
    text = re.sub(r"\n{3,}", "\n\n", text.strip())
    return text[:limit]


# ---------------------------------------------------------------- 附件预览

#: 预览最多渲染多少字 / 多少行。预览是"看一眼确认带对了东西"，不是阅读器——
#: 一份 300 页的规范全渲染出来，浏览器卡住的代价远大于多看两页的收益。
PREVIEW_TEXT_CHARS = 40_000
PREVIEW_TABLE_ROWS = 200
PREVIEW_TABLE_COLS = 40

_MARKDOWN_EXTS = frozenset({".md", ".markdown"})
_SHEET_EXTS = frozenset({".xlsx", ".xlsm"})
_CSV_EXTS = frozenset({".csv", ".tsv"})


def _sheet_rows(data: bytes) -> tuple[list[list[str]], str]:
    """xlsx → 二维字符串表 + 工作表名。只读第一个工作表。"""
    from io import BytesIO

    from openpyxl import load_workbook

    # read_only + data_only：只要值不要公式与样式，几十 MB 的表也不会把内存吃光
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows: list[list[str]] = []
        for row in ws.iter_rows(max_row=PREVIEW_TABLE_ROWS, max_col=PREVIEW_TABLE_COLS):
            rows.append(["" if c.value is None else str(c.value) for c in row])
        # 末尾的整行空行去掉：openpyxl 常把带过格式的空行也算进 dimensions
        while rows and not any(cell.strip() for cell in rows[-1]):
            rows.pop()
        # 尾部空列同理，而且必须裁：`max_col` 让 openpyxl 把每行补齐到 40 格，
        # 一张两列的表会渲染出 38 个空列，表格宽得要横向滚动才看得到内容
        width = 0
        for row in rows:
            for i, cell in enumerate(row):
                if cell.strip():
                    width = max(width, i + 1)
        rows = [row[:width] for row in rows]
        return rows, str(ws.title or "")
    finally:
        wb.close()


def _csv_rows(text: str, delimiter: str) -> list[list[str]]:
    import csv
    from io import StringIO

    out: list[list[str]] = []
    for i, row in enumerate(csv.reader(StringIO(text), delimiter=delimiter)):
        if i >= PREVIEW_TABLE_ROWS:
            break
        out.append([str(c) for c in row[:PREVIEW_TABLE_COLS]])
    return out


def document_preview(name: str, data: bytes) -> dict:
    """附件预览的结构化结果。

    `kind` 决定前端怎么渲染，四种：

    | kind | 内容 | 前端 |
    | --- | --- | --- |
    | `markdown` | `text` | react-markdown（本仓已有依赖） |
    | `text` | `text` | 等宽 `<pre>` |
    | `table` | `rows` + `sheet` | 表格 |
    | `binary` | 只有 `name` | 一句"看不了，下载吧" |

    **pdf 不在这里**：它由浏览器内置的 PDF 阅读器直接渲染（`<iframe>` 指向
    原文件），比把版式拆成纯文本再拼回去忠实得多，也不用引 pdf.js。
    """
    ext = ("." + name.rsplit(".", 1)[-1].lower()) if "." in name else ""
    try:
        if ext in _SHEET_EXTS:
            rows, sheet = _sheet_rows(data)
            return {"kind": "table", "rows": rows, "sheet": sheet}
        if ext in _CSV_EXTS:
            text = decode_text(data)
            return {
                "kind": "table",
                "rows": _csv_rows(text, "\t" if ext == ".tsv" else ","),
                "sheet": "",
            }
        if ext in _MARKDOWN_EXTS:
            return {"kind": "markdown", "text": decode_text(data)[:PREVIEW_TEXT_CHARS]}
        if ext in _TEXT_EXTS:
            return {"kind": "text", "text": decode_text(data)[:PREVIEW_TEXT_CHARS]}
    except Exception:  # noqa: BLE001 —— 预览失败退化成"下载吧"，不该 500
        return {"kind": "binary"}
    return {"kind": "binary"}


async def asset_preview(session: AsyncSession, asset_id: int) -> dict:
    """按 id 取媒体资产的预览。资产不存在时抛 KeyError 由路由层转 404。"""
    row = await session.get(StudioMediaAsset, asset_id)
    if row is None:
        raise KeyError(asset_id)
    try:
        data = await get_storage().read(row.storage_key)
    except Exception:  # noqa: BLE001 —— 存储读不到就当"看不了"
        return {"kind": "binary", "name": row.name}
    out = document_preview(row.name, data)
    out["name"] = row.name
    return out


async def asset_text(session: AsyncSession, asset_id: int) -> tuple[str, str]:
    """按 id 取媒体资产的 `(文件名, 正文)`。资产不存在时返回两个空串。"""
    row = await session.get(StudioMediaAsset, asset_id)
    if row is None:
        return "", ""
    try:
        data = await get_storage().read(row.storage_key)
    except Exception:  # noqa: BLE001 —— 存储读不到时退化成"只知道文件名"
        return row.name, ""
    return row.name, document_text(row.name, data)


def kind_for_upload(name: str, mime: str) -> str:
    """按 MIME 优先、扩展名兜底识别画布上传素材。"""
    normalized_mime = mime.strip().lower()
    if normalized_mime.startswith("video/"):
        return "video"
    if normalized_mime.startswith("audio/"):
        return "audio"
    extension = mimetypes.guess_extension(normalized_mime) or ""
    if not extension:
        extension = "." + name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if extension in VIDEO_EXTENSIONS:
        return "video"
    if extension in AUDIO_EXTENSIONS:
        return "audio"
    return "file"


def asset_view(row: StudioMediaAsset) -> dict[str, Any]:
    return {
        "id": row.id,
        "kind": row.kind,
        "name": row.name,
        "mime": row.mime,
        "bytes": row.bytes,
        "width": row.width,
        "height": row.height,
        "duration_ms": row.duration_ms,
        "source_task_id": row.source_task_id,
        "source_url": row.source_url,
        "details": row.details,
        "parent_id": row.parent_id,
        "group_id": row.group_id,
        "status": row.status,
        "favorite": row.favorite,
        "url": f"/api/studio/media-assets/{row.id}/content",
        "poster_url": (
            f"/api/studio/media-assets/{row.id}/poster"
            if row.poster_key
            else None
        ),
        "created_at": row.created_at.isoformat() if row.created_at else None,
    }


async def ingest_one(
    session: AsyncSession,
    data: bytes,
    *,
    kind: str,
    name: str,
    mime: str,
    source_task_id: str | None = None,
    source_url: str | None = None,
    details: dict[str, Any] | None = None,
    width: int | None = None,
    height: int | None = None,
    duration_ms: int | None = None,
    parent_id: int | None = None,
) -> StudioMediaAsset:
    normalized_kind = kind.strip().lower()
    if normalized_kind not in MEDIA_KINDS:
        raise StudioMediaAssetError(f"不支持的媒体资产类型：{kind}")
    if not data:
        raise StudioMediaAssetError("媒体资产内容为空")
    digest = hashlib.sha256(data).hexdigest()
    existing = (
        await session.execute(
            select(StudioMediaAsset).where(StudioMediaAsset.sha256 == digest)
        )
    ).scalar_one_or_none()
    if existing is not None:
        return existing

    clean_name = re.sub(r"[^a-zA-Z0-9._-]+", "-", name).strip(".-")
    extension = mimetypes.guess_extension(mime) or ""
    if not clean_name:
        clean_name = f"{normalized_kind}{extension}"
    if extension and "." not in clean_name:
        clean_name += extension
    key = (
        f"studio-media/{normalized_kind}/{source_task_id or 'import'}/"
        f"{uuid.uuid4().hex[:10]}-{clean_name}"
    )
    storage = get_storage()
    await storage.write(key, data)
    row = StudioMediaAsset(
        kind=normalized_kind,
        name=name.strip() or clean_name,
        mime=mime.strip() or "application/octet-stream",
        sha256=digest,
        storage_key=key,
        bytes=len(data),
        width=width,
        height=height,
        duration_ms=duration_ms,
        source_task_id=source_task_id,
        source_url=source_url,
        details=details,
        parent_id=parent_id,
    )
    session.add(row)
    try:
        await session.flush()
    except Exception:
        await storage.delete(key)
        raise
    return row
